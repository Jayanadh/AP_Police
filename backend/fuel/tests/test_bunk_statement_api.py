"""A bunk statement: every fill made at one pump in a period, with totals. Nothing is submitted or disputed."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from common.months import IST
from fleet.models import FuelType
from fuel.models import RequestStatus
from testing.factories import (
    DriverFactory,
    FuelRequestFactory,
    MTOFactory,
    OfficerFactory,
    PTOFactory,
    PumpStaffFactory,
    UnitFactory,
    VehicleFactory,
    filled_request,
    police_pump,
    tieup_pump,
)

pytestmark = pytest.mark.django_db

URL = "/api/fuel/bunk-statement/"
EXPORT = "/api/fuel/bunk-statement/export/"
SEPTEMBER = {"from": "2026-09-01", "to": "2026-09-30"}


def at(day, hour=10):
    return datetime(2026, 9, day, hour, 0, tzinfo=IST)


@pytest.fixture
def unit():
    return UnitFactory(name="SPSR Nellore MTO")


@pytest.fixture
def bunk(unit):
    pump = tieup_pump(unit)
    pump.name = "Sri Venkateswara Fuels"
    pump.save()
    return pump


@pytest.fixture
def fills(unit, bunk):
    """Two September fills at the bunk, one by a vehicle of another office, and fills that do not belong."""
    jeep = VehicleFactory(unit=unit, registration_number="AP39PA1001", fuel_type=FuelType.DIESEL)
    guntur = UnitFactory(name="Guntur MTO")
    car = VehicleFactory(unit=guntur, registration_number="AP07PB2001", fuel_type=FuelType.PETROL)
    ravi = DriverFactory(unit=unit, full_name="Ravi Kumar")
    later = filled_request(jeep, Decimal("40"), at(20), driver=ravi, pump=bunk)
    earlier = filled_request(car, Decimal("25.5"), at(3), driver=DriverFactory(full_name="G. Prasad"), pump=bunk)
    filled_request(jeep, Decimal("30"), datetime(2026, 10, 1, 9, 0, tzinfo=IST), pump=bunk)  # October
    filled_request(jeep, Decimal("12"), at(5), pump=tieup_pump(unit))  # another bunk
    FuelRequestFactory(vehicle=jeep, pump=bunk, status=RequestStatus.ISSUED)  # not filled
    return earlier, later


def test_the_mto_sees_every_fill_at_one_of_the_offices_bunks_in_a_period(api, unit, bunk, fills):
    earlier, later = fills
    api.force_login(MTOFactory(unit=unit))

    response = api.get(URL, {"pump": bunk.id, **SEPTEMBER})

    assert response.status_code == 200
    body = response.json()
    assert body["from"] == "2026-09-01"
    assert body["to"] == "2026-09-30"
    assert body["label"] == "September 2026"
    assert body["pump"] == {
        "id": bunk.id, "name": "Sri Venkateswara Fuels", "kind": "TIE_UP", "kind_label": "Tie-up bunk",
        "unit_name": "SPSR Nellore MTO",
    }
    assert (body["fills"], body["litres"]) == (2, "65.50")
    assert (body["petrol_litres"], body["diesel_litres"]) == ("25.50", "40.00")
    assert [row["id"] for row in body["rows"]] == [earlier.id, later.id]  # in the order they were filled
    assert body["rows"][0] == {
        "id": earlier.id,
        "filled_at": "2026-09-03T10:00:00+05:30",
        "registration_number": "AP07PB2001",
        "unit_name": "Guntur MTO",
        "driver_name": "G. Prasad",
        "fuel_type": "PETROL",
        "litres": "25.50",
    }


def test_a_police_pump_has_a_statement_too(api, unit):
    pump = police_pump(unit, diesel=Decimal("500"))
    filled_request(VehicleFactory(unit=unit), Decimal("10"), at(4), pump=pump)
    api.force_login(MTOFactory(unit=unit))

    assert api.get(URL, {"pump": pump.id, **SEPTEMBER}).json()["litres"] == "10.00"


def test_the_pto_sees_any_offices_pump(api, bunk, fills):
    api.force_login(PTOFactory())

    assert api.get(URL, {"pump": bunk.id, **SEPTEMBER}).json()["fills"] == 2


def test_another_offices_mto_cannot_see_the_bunk(api, bunk, fills):
    api.force_login(MTOFactory())

    assert api.get(URL, {"pump": bunk.id, **SEPTEMBER}).status_code == 404


def test_pump_staff_see_their_own_pump_without_naming_it(api, unit, bunk, fills):
    api.force_login(PumpStaffFactory(pump=bunk, unit=unit))

    body = api.get(URL, SEPTEMBER).json()

    assert body["pump"]["id"] == bunk.id
    assert body["fills"] == 2


def test_pump_staff_cannot_see_another_pump(api, unit, bunk):
    other = tieup_pump(unit)
    api.force_login(PumpStaffFactory(pump=bunk, unit=unit))

    assert api.get(URL, {"pump": other.id, **SEPTEMBER}).status_code == 404


def test_the_mto_and_pto_must_pick_a_pump(api, unit):
    api.force_login(MTOFactory(unit=unit))

    assert api.get(URL, SEPTEMBER).json() == {"detail": "Pick a bunk."}
    assert api.get(URL, {"pump": "abc", **SEPTEMBER}).json() == {"detail": "Use the pump's id."}


@pytest.mark.parametrize("maker", [OfficerFactory, DriverFactory])
def test_officers_and_drivers_have_no_bunk_statements(api, bunk, maker):
    api.force_login(maker())

    assert api.get(URL, {"pump": bunk.id}).status_code == 403
    assert api.get(EXPORT, {"pump": bunk.id}).status_code == 403


def test_a_period_too_long_is_refused(api, unit, bunk):
    api.force_login(MTOFactory(unit=unit))

    response = api.get(URL, {"pump": bunk.id, "from": "2025-01-01", "to": "2026-09-30"})

    assert response.status_code == 400
    assert response.json() == {"detail": "Pick a period of at most a year."}


def test_the_statement_downloads_as_excel_with_every_fill_and_the_totals(api, unit, bunk, fills):
    api.force_login(MTOFactory(unit=unit))

    response = api.get(EXPORT, {"pump": bunk.id, **SEPTEMBER})

    assert response.status_code == 200
    assert response["Content-Disposition"] == (
        'attachment; filename="bunk-statement-sri-venkateswara-fuels-2026-09-01-to-2026-09-30.xlsx"'
    )
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Fills", "Totals"]
    assert [[cell.value for cell in row] for row in book["Fills"].iter_rows()] == [
        ["Date", "Vehicle", "Office", "Driver", "Fuel", "Litres"],
        [datetime(2026, 9, 3, 10, 0), "AP07PB2001", "Guntur MTO", "G. Prasad", "Petrol", 25.5],
        [datetime(2026, 9, 20, 10, 0), "AP39PA1001", "SPSR Nellore MTO", "Ravi Kumar", "Diesel", 40],
    ]
    assert [[cell.value for cell in row] for row in book["Totals"].iter_rows()] == [
        ["Bunk", "Office", "Period", "Fills", "Petrol (L)", "Diesel (L)", "Total (L)"],
        ["Sri Venkateswara Fuels", "SPSR Nellore MTO", "September 2026", 2, 25.5, 40, 65.5],
    ]


def test_the_old_statements_address_is_gone(api, unit):
    api.force_login(MTOFactory(unit=unit))

    assert api.get("/api/fuel/statements/").status_code == 404


# --- the pumps a statement can be drawn up for -----------------------------------------------------------------

PUMPS = "/api/fuel/bunk-statement/pumps/"


def test_the_mto_picks_from_the_offices_pumps_closed_ones_too(api, unit, bunk):
    closed = police_pump(unit)
    closed.name, closed.is_active = "Old Police Pump", False
    closed.save()
    tieup_pump(UnitFactory())  # another office's
    api.force_login(MTOFactory(unit=unit))

    rows = api.get(PUMPS).json()

    assert rows == [
        {
            "id": closed.id, "name": "Old Police Pump", "kind": "POLICE", "kind_label": "Police pump",
            "unit": unit.id, "unit_name": "SPSR Nellore MTO", "is_active": False,
        },
        {
            "id": bunk.id, "name": "Sri Venkateswara Fuels", "kind": "TIE_UP", "kind_label": "Tie-up bunk",
            "unit": unit.id, "unit_name": "SPSR Nellore MTO", "is_active": True,
        },
    ]


def test_the_pto_picks_from_every_offices_pumps_or_one_offices(api, unit, bunk):
    other = tieup_pump(UnitFactory(name="Guntur MTO"))
    api.force_login(PTOFactory())

    assert {row["id"] for row in api.get(PUMPS).json()} == {bunk.id, other.id}
    assert [row["id"] for row in api.get(PUMPS, {"unit": unit.id}).json()] == [bunk.id]


@pytest.mark.parametrize("maker", [OfficerFactory, DriverFactory])
def test_officers_and_drivers_have_no_pumps_to_pick(api, maker):
    api.force_login(maker())

    assert api.get(PUMPS).status_code == 403
