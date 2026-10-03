"""/api/monitor/…: the PTO's read-only view of every office's vehicles, officers and drivers."""
from datetime import timedelta
from decimal import Decimal
from types import SimpleNamespace

import pytest
from django.utils import timezone

from accounts.models import UserStatus
from common.months import current_month, month_bounds, today_ist
from fleet import odometer, services
from fleet.models import VehicleStatus
from fuel.models import FuelRequest
from testing.factories import (
    DriverFactory,
    FuelGrantFactory,
    MTOFactory,
    OdometerReadingFactory,
    OfficerFactory,
    PTOFactory,
    PumpStaffFactory,
    UnitFactory,
    VehicleFactory,
    filled_request,
)

pytestmark = pytest.mark.django_db

VEHICLES = "/api/monitor/vehicles/"
OFFICERS = "/api/monitor/officers/"
DRIVERS = "/api/monitor/drivers/"


@pytest.fixture
def world():
    """Nellore has a Bolero (driver and officer linked, a 30 L grant, 40 L drawn this month, an odometer reading past
    its service interval) and a paused Innova; Guntur has a Swift."""
    nellore = UnitFactory(name="Nellore MTO")
    guntur = UnitFactory(name="Guntur MTO")
    mto = MTOFactory(unit=nellore)
    bolero = VehicleFactory(
        unit=nellore, registration_number="AP39PA0001", make="Mahindra", model="Bolero",
        monthly_fuel_limit_litres=150, service_interval_km=1000,
    )
    innova = VehicleFactory(unit=nellore, registration_number="AP39PA0002", status=VehicleStatus.PAUSED)
    swift = VehicleFactory(unit=guntur, registration_number="AP07PB0001")
    driver = DriverFactory(
        unit=nellore, full_name="Ravi Kumar", emp_id="AP4001", mobile="9876543210", licence_number="AP0123"
    )
    officer = OfficerFactory(unit=nellore, full_name="S. Venkata Rao", emp_id="AP3001")
    services.assign_person(bolero, driver, mto)
    services.assign_person(bolero, officer, mto)
    FuelGrantFactory(vehicle=bolero, litres=Decimal("30"))
    filled_request(bolero, Decimal("40"), timezone.now() - timedelta(minutes=5), driver=driver)
    reading = OdometerReadingFactory(
        vehicle=bolero, week_of=odometer.week_of(today_ist()), reading_km=1500, recorded_by=driver
    )
    return SimpleNamespace(
        nellore=nellore, guntur=guntur, mto=mto, bolero=bolero, innova=innova, swift=swift, driver=driver,
        officer=officer, reading=reading,
    )


def read(api, url, **params):
    api.force_login(PTOFactory())
    response = api.get(url, params)
    assert response.status_code == 200, response.content
    return response.json()


def registrations(body):
    return [row["registration_number"] for row in body["results"]]


# --- vehicles --------------------------------------------------------------------------------------------------


def test_the_pto_sees_every_offices_vehicles_by_office_then_registration(api, world):
    body = read(api, VEHICLES)

    assert (body["count"], body["page"], body["pages"], body["page_size"]) == (3, 1, 1, 50)
    assert registrations(body) == ["AP07PB0001", "AP39PA0001", "AP39PA0002"]


def test_a_vehicle_row_has_what_monitoring_needs(api, world):
    row = read(api, VEHICLES, search="AP39PA0001")["results"][0]

    assert row == {
        "id": world.bolero.pk,
        "registration_number": "AP39PA0001",
        "vehicle_type": world.bolero.vehicle_type,
        "make": "Mahindra",
        "model": "Bolero",
        "fuel_type": world.bolero.fuel_type,
        "unit": world.nellore.pk,
        "unit_name": "Nellore MTO",
        "status": "ACTIVE",
        "status_label": "Active",
        "driver": {"id": world.driver.pk, "full_name": "Ravi Kumar", "emp_id": "AP4001", "mobile": "9876543210"},
        "officer": {"id": world.officer.pk, "full_name": "S. Venkata Rao", "emp_id": "AP3001", "mobile": ""},
        "limit_litres": "180.00",
        "used_litres": "40.00",
        "remaining_litres": "140.00",
        "last_odometer_km": 1500,
        "last_odometer_week": world.reading.week_of.isoformat(),
        "service_due": True,
    }


def test_a_vehicle_with_nobody_linked_and_no_readings(api, world):
    row = read(api, VEHICLES, search="AP07PB0001")["results"][0]

    assert (row["driver"], row["officer"]) == (None, None)
    assert (row["last_odometer_km"], row["last_odometer_week"]) == (None, None)
    assert (row["used_litres"], row["service_due"]) == ("0.00", False)


def test_a_terminated_vehicle_is_never_due_for_service(api, world):
    world.bolero.status = VehicleStatus.TERMINATED
    world.bolero.save(update_fields=["status"])

    assert read(api, VEHICLES, search="AP39PA0001")["results"][0]["service_due"] is False


def test_vehicles_can_be_narrowed_by_office_status_and_registration(api, world):
    assert registrations(read(api, VEHICLES, unit=str(world.guntur.pk))) == ["AP07PB0001"]
    assert registrations(read(api, VEHICLES, status="PAUSED")) == ["AP39PA0002"]
    assert registrations(read(api, VEHICLES, search="ap39 pa 0002")) == ["AP39PA0002"]
    assert registrations(read(api, VEHICLES, unit=str(world.guntur.pk), status="PAUSED")) == []


def test_vehicles_come_fifty_to_a_page(api, world):
    VehicleFactory.create_batch(48, unit=world.guntur)

    first = read(api, VEHICLES)
    second = read(api, VEHICLES, page="2")

    assert (first["count"], first["pages"], len(first["results"])) == (51, 2, 50)
    assert (second["page"], len(second["results"])) == (2, 1)
    assert registrations(second) == ["AP39PA0002"]


def test_a_page_of_vehicles_takes_the_same_queries_however_many_there_are(api, world, django_assert_max_num_queries):
    api.force_login(PTOFactory())
    with django_assert_max_num_queries(8):
        api.get(VEHICLES)
    for vehicle in VehicleFactory.create_batch(10, unit=world.nellore):
        services.assign_person(vehicle, DriverFactory(unit=world.nellore), world.mto)
    with django_assert_max_num_queries(8):
        assert len(api.get(VEHICLES).json()["results"]) == 13


# --- officers and drivers --------------------------------------------------------------------------------------


def test_the_pto_sees_every_offices_officers_with_their_vehicles(api, world):
    OfficerFactory(unit=world.guntur, full_name="A. Prasad", emp_id="AP3101")

    body = read(api, OFFICERS)

    assert [row["full_name"] for row in body["results"]] == ["A. Prasad", "S. Venkata Rao"]
    row = body["results"][1]
    assert row == {
        "id": world.officer.pk,
        "full_name": "S. Venkata Rao",
        "emp_id": "AP3001",
        "designation_name": "Inspector of Police",
        "unit": world.nellore.pk,
        "unit_name": "Nellore MTO",
        "status": "ACTIVE",
        "status_label": "Active",
        "mobile": "",
        "current_vehicles": [{"id": world.bolero.pk, "registration_number": "AP39PA0001"}],
    }


def test_drivers_also_show_their_licence_fuel_this_month_and_overdue_duty_particulars(api, world):
    overdue_at = timezone.now() - timedelta(hours=50)
    filled_request(world.bolero, Decimal("10"), overdue_at, driver=world.driver)
    month_start, _ = month_bounds(current_month())
    litres_this_month = Decimal("40") + (Decimal("10") if overdue_at >= month_start else Decimal("0"))

    row = read(api, DRIVERS, search="AP4001")["results"][0]

    assert row["licence_number"] == "AP0123"
    assert row["licence_valid_till"] is None
    assert row["current_vehicles"] == [{"id": world.bolero.pk, "registration_number": "AP39PA0001"}]
    assert row["overdue_duty"] == 1
    assert row["litres_this_month"] == f"{litres_this_month:.2f}"
    assert row["fills_this_month"] == (2 if overdue_at >= month_start else 1)


def test_written_duty_particulars_are_not_overdue(api, world):
    late = filled_request(world.bolero, Decimal("10"), timezone.now() - timedelta(hours=50), driver=world.driver)
    FuelRequest.objects.filter(pk=late.pk).update(duty_particulars="Patrol", duty_submitted_at=timezone.now())

    assert read(api, DRIVERS, search="AP4001")["results"][0]["overdue_duty"] == 0


def test_people_can_be_narrowed_by_office_status_and_name_or_emp_id(api, world):
    paused = DriverFactory(unit=world.guntur, full_name="M. Rao", status=UserStatus.PAUSED)

    def names(**params):
        return [row["full_name"] for row in read(api, DRIVERS, **params)["results"]]

    assert names(unit=str(world.guntur.pk)) == ["M. Rao"]
    assert names(status="PAUSED") == ["M. Rao"]
    assert names(search="ravi") == ["Ravi Kumar"]
    assert names(search=paused.emp_id.lower()) == ["M. Rao"]


def test_the_lists_hold_only_their_own_role(api, world):
    assert [row["full_name"] for row in read(api, DRIVERS)["results"]] == ["Ravi Kumar"]
    assert [row["full_name"] for row in read(api, OFFICERS)["results"]] == ["S. Venkata Rao"]


# --- who may look ----------------------------------------------------------------------------------------------


@pytest.mark.parametrize("url", [VEHICLES, OFFICERS, DRIVERS])
@pytest.mark.parametrize("make_user", [MTOFactory, OfficerFactory, DriverFactory, PumpStaffFactory])
def test_only_the_pto_can_monitor(api, url, make_user):
    api.force_login(make_user())

    assert api.get(url).status_code == 403


@pytest.mark.parametrize("url", [VEHICLES, OFFICERS, DRIVERS])
def test_the_lists_are_read_only(api, world, url):
    api.force_login(PTOFactory())

    assert api.post(url, {}).status_code == 405


@pytest.mark.parametrize("url", [VEHICLES, OFFICERS, DRIVERS])
def test_the_office_filter_must_be_an_id(api, url):
    api.force_login(PTOFactory())

    response = api.get(url, {"unit": "Nellore"})

    assert response.status_code == 400
    assert response.json() == {"detail": "Use the office's id."}


# --- the lists as Excel files ---------------------------------------------------------------------------------


def download(api, url, **params):
    from io import BytesIO

    from openpyxl import load_workbook

    api.force_login(PTOFactory())
    response = api.get(f"{url}export/", params)
    assert response.status_code == 200, response.content
    sheet = load_workbook(BytesIO(response.content)).worksheets[0]
    return response, [[cell.value for cell in row] for row in sheet.iter_rows()]


def test_the_vehicles_download_with_the_filters_applied_and_every_page(api, world):
    response, rows = download(api, VEHICLES, unit=world.nellore.id)

    assert response["Content-Disposition"] == 'attachment; filename="vehicles.xlsx"'
    assert rows[0] == [
        "Registration", "Vehicle", "Fuel", "Office", "Status", "Driver", "Officer", "Limit (L)", "Used (L)",
        "Left (L)", "Last odometer (km)", "Service due",
    ]
    assert rows[1] == [
        "AP39PA0001", "Mahindra Bolero", "Diesel", "Nellore MTO", "Active", "Ravi Kumar", "S. Venkata Rao", 180,
        40, 140, 1500, "Yes",
    ]
    assert [row[0] for row in rows[1:]] == ["AP39PA0001", "AP39PA0002"]


def test_the_drivers_download_with_their_licence_and_this_months_fuel(api, world):
    response, rows = download(api, DRIVERS)

    assert response["Content-Disposition"] == 'attachment; filename="drivers.xlsx"'
    assert rows[0] == [
        "Name", "Emp ID", "Designation", "Office", "Status", "Mobile", "Vehicles", "Licence", "Licence valid till",
        "Fills this month", "Litres this month", "Overdue duty particulars",
    ]
    ravi = next(row for row in rows if row[0] == "Ravi Kumar")
    assert ravi[1] == "AP4001"
    assert ravi[5:8] == ["9876543210", "AP39PA0001", "AP0123"]
    assert ravi[9:] == [1, 40, 0]


def test_the_officers_download(api, world):
    response, rows = download(api, OFFICERS, search="Venkata")

    assert response["Content-Disposition"] == 'attachment; filename="officers.xlsx"'
    assert rows[0] == ["Name", "Emp ID", "Designation", "Office", "Status", "Mobile", "Vehicles"]
    assert [row[0] for row in rows[1:]] == ["S. Venkata Rao"]


@pytest.mark.parametrize("url", [VEHICLES, OFFICERS, DRIVERS])
def test_only_the_pto_downloads_the_lists(api, world, url):
    api.force_login(world.mto)

    assert api.get(f"{url}export/").status_code == 403
