from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from common.excel import Sheet, xlsx_response
from common.months import IST


def read(response):
    return load_workbook(BytesIO(response.content))


def test_the_response_is_an_excel_download_with_the_given_name():
    response = xlsx_response("fuel-statement-2026-09.xlsx", [Sheet("Fills", ["Vehicle"], [])])

    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response["Content-Disposition"] == 'attachment; filename="fuel-statement-2026-09.xlsx"'


def test_each_sheet_has_a_bold_frozen_header_and_its_rows():
    sheets = [
        Sheet("Fills", ["Vehicle", "Litres"], [["AP39PA1001", Decimal("40.50")], ["AP39PA1002", Decimal("12")]]),
        Sheet("Totals", ["Fuel", "Litres"], [["Diesel", Decimal("52.50")]]),
    ]

    book = read(xlsx_response("x.xlsx", sheets))

    assert book.sheetnames == ["Fills", "Totals"]
    fills = book["Fills"]
    assert [[cell.value for cell in row] for row in fills.iter_rows()] == [
        ["Vehicle", "Litres"],
        ["AP39PA1001", 40.5],
        ["AP39PA1002", 12],
    ]
    assert fills["A1"].font.bold
    assert fills.freeze_panes == "A2"
    assert fills["B2"].number_format == "#,##0.00"


def test_times_are_written_in_india_time_and_dates_as_dates():
    filled = datetime(2026, 9, 30, 20, 0, tzinfo=IST).astimezone(tz=None)  # whatever zone the server runs in

    sheet = read(xlsx_response("x.xlsx", [Sheet("Fills", ["When", "Day"], [[filled, date(2026, 9, 1)]])]))["Fills"]

    assert sheet["A2"].value == datetime(2026, 9, 30, 20, 0)
    assert sheet["A2"].number_format == "dd mmm yyyy hh:mm"
    assert sheet["B2"].value == datetime(2026, 9, 1)
    assert sheet["B2"].number_format == "dd mmm yyyy"


def test_typed_text_never_becomes_a_formula():
    """Duty particulars or remarks such as `=HYPERLINK(...)` must open as text, not run as a formula."""
    risky = ['=HYPERLINK("http://example.com","x")', "+1+1", "-2+3", "@SUM(A1)"]

    sheet = read(xlsx_response("x.xlsx", [Sheet("Fills", ["Note"], [[text] for text in risky])]))["Fills"]

    assert [sheet.cell(row=row, column=1).value for row in range(2, 6)] == risky
    assert {sheet.cell(row=row, column=1).data_type for row in range(2, 6)} == {"s"}


def test_empty_values_stay_empty_and_long_sheet_names_are_cut_to_excels_limit():
    book = read(xlsx_response("x.xlsx", [Sheet("Vehicles of SPSR Nellore MTO: September", ["A", "B"], [[None, "x"]])]))

    sheet = book.worksheets[0]
    assert sheet.title == "Vehicles of SPSR Nellore MTO Se"
    assert sheet["A2"].value is None
