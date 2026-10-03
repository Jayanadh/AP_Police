"""Excel downloads: each report hands over its rows, and this writes them as an .xlsx file.

Text is always written as text. openpyxl would otherwise store a value that starts with "=" as a formula, so a driver's
duty particulars could run as a formula on the MTO's computer when the file is opened.
"""
import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from common.months import IST

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LITRES_FORMAT = "#,##0.00"
DATE_FORMAT = "dd mmm yyyy"
DATETIME_FORMAT = "dd mmm yyyy hh:mm"
MAX_TITLE = 31  # Excel's limit on a sheet's name
MAX_WIDTH = 60
_NOT_IN_TITLES = re.compile(r"[\[\]:*?/\\]")


@dataclass(frozen=True)
class Sheet:
    title: str
    headers: Sequence[str]
    rows: Iterable[Sequence[object]]


def xlsx_response(filename: str, sheets: Sequence[Sheet]) -> HttpResponse:
    """`sheets` as an .xlsx download called `filename`."""
    book = Workbook()
    book.remove(book.active)
    for sheet in sheets:
        _write(book.create_sheet(_NOT_IN_TITLES.sub("", sheet.title)[:MAX_TITLE]), sheet)
    content = BytesIO()
    book.save(content)
    response = HttpResponse(content.getvalue(), content_type=CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _write(worksheet, sheet: Sheet) -> None:
    worksheet.append(list(sheet.headers))
    widths = [len(header) for header in sheet.headers]
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for row_number, values in enumerate(sheet.rows, start=2):
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column)
            _set(cell, value)
            widths[column - 1] = max(widths[column - 1], len(_shown(value)))
    worksheet.freeze_panes = "A2"
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = min(width + 2, MAX_WIDTH)


def _set(cell: Cell, value: object) -> None:
    if isinstance(value, datetime):
        cell.value = value.astimezone(IST).replace(tzinfo=None) if value.tzinfo else value
        cell.number_format = DATETIME_FORMAT
    elif isinstance(value, date):
        cell.value = datetime(value.year, value.month, value.day)
        cell.number_format = DATE_FORMAT
    elif isinstance(value, Decimal):
        cell.value = value
        cell.number_format = LITRES_FORMAT
    elif isinstance(value, str):
        cell.value = value
        cell.data_type = "s"  # never a formula, whatever it starts with
    else:
        cell.value = value


def _shown(value: object) -> str:
    """Roughly what the cell shows, to size its column."""
    if isinstance(value, datetime):
        return "30 Sep 2026 20:00"
    if isinstance(value, date):
        return "30 Sep 2026"
    return "" if value is None else str(value)
