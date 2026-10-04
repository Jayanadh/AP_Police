import pytest

from accounts.models import UserStatus
from fleet import services
from fleet.models import VehicleAssignment, VehicleStatus
from testing.factories import (
    DriverFactory,
    MTOFactory,
    OdometerReadingFactory,
    OfficerFactory,
    PTOFactory,
    ServiceRecordFactory,
    UnitFactory,
    VehicleFactory,
)

pytestmark = pytest.mark.django_db


@pytest.fixture
def mto():
    return MTOFactory(full_name="Sri K. Ramesh", mobile="9876500001")


@pytest.fixture
def mto_api(api, mto):
    api.force_login(mto)
    return api


@pytest.fixture
def vehicle(mto):
    return VehicleFactory(unit=mto.unit)


def links_url(vehicle):
    return f"/api/vehicles/{vehicle.id}/assignments/"


# --- linking people to a vehicle -------------------------------------------------------------


def test_mto_links_an_officer_then_a_driver_and_the_vehicle_shows_both(mto_api, mto, vehicle):
    officer = OfficerFactory(unit=mto.unit, full_name="Insp. Rao", emp_id="EMP00901")
    driver = DriverFactory(unit=mto.unit)

    officer_response = mto_api.post(links_url(vehicle), {"person": officer.id})
    driver_response = mto_api.post(links_url(vehicle), {"person": driver.id})

    assert officer_response.status_code == 201
    assert officer_response.json() == {
        "id": officer_response.json()["id"],
        "vehicle": vehicle.id,
        "person": officer.id,
        "person_name": "Insp. Rao",
        "person_emp_id": "EMP00901",
        "kind": "OFFICER",
        "started_at": officer_response.json()["started_at"],
        "ended_at": None,
        "assigned_by_name": mto.full_name,
        "ended_by_name": None,
    }
    assert driver_response.status_code == 201
    assert driver_response.json()["kind"] == "DRIVER"

    detail = mto_api.get(f"/api/vehicles/{vehicle.id}/").json()
    assert detail["current_officer"]["id"] == officer.id
    assert detail["current_officer"]["assignment_id"] == officer_response.json()["id"]
    assert detail["current_driver"]["id"] == driver.id
    assert VehicleAssignment.objects.get(pk=driver_response.json()["id"]).assigned_by == mto


def test_a_second_officer_is_refused_with_the_service_message(mto_api, mto, vehicle):
    services.assign_person(vehicle, OfficerFactory(unit=mto.unit), mto)
    second = OfficerFactory(unit=mto.unit)

    response = mto_api.post(links_url(vehicle), {"person": second.id})

    assert response.status_code == 400
    assert response.json() == {"detail": "This vehicle already has an officer. End that link first."}
    assert VehicleAssignment.objects.filter(vehicle=vehicle).count() == 1


def test_a_second_driver_is_refused_with_the_service_message(mto_api, mto, vehicle):
    services.assign_person(vehicle, DriverFactory(unit=mto.unit), mto)
    second = DriverFactory(unit=mto.unit)

    response = mto_api.post(links_url(vehicle), {"person": second.id})

    assert response.status_code == 400
    assert response.json() == {"detail": "This vehicle already has a driver. End that link first."}
    assert VehicleAssignment.objects.filter(vehicle=vehicle).count() == 1


def test_a_driver_already_on_another_vehicle_is_refused(mto_api, mto, vehicle):
    driver = DriverFactory(unit=mto.unit)
    services.assign_person(VehicleFactory(unit=mto.unit), driver, mto)

    response = mto_api.post(links_url(vehicle), {"person": driver.id})

    assert response.status_code == 400
    assert response.json() == {"detail": "This driver is already linked to another vehicle."}


def test_a_person_from_another_unit_is_refused_on_the_person_field(mto_api, vehicle):
    outsider = DriverFactory(unit=UnitFactory())

    response = mto_api.post(links_url(vehicle), {"person": outsider.id})

    assert response.status_code == 400
    assert "person" in response.json()
    assert VehicleAssignment.objects.count() == 0


@pytest.mark.parametrize("payload", [{}, {"person": None}, {"person": 999999}, {"person": "abc"}])
def test_a_missing_or_unknown_person_is_refused_on_the_person_field(mto_api, vehicle, payload):
    response = mto_api.post(links_url(vehicle), payload)

    assert response.status_code == 400
    assert "person" in response.json()


def test_the_mto_login_cannot_be_linked(mto_api, mto, vehicle):
    response = mto_api.post(links_url(vehicle), {"person": mto.id})

    assert response.status_code == 400
    assert response.json() == {"detail": "Only officers and drivers can be linked to vehicles."}


def test_a_paused_person_cannot_be_linked(mto_api, mto, vehicle):
    driver = DriverFactory(unit=mto.unit, status=UserStatus.PAUSED)

    response = mto_api.post(links_url(vehicle), {"person": driver.id})

    assert response.status_code == 400
    assert response.json() == {"detail": "Only active officers and drivers can be linked."}


def test_a_terminated_vehicle_cannot_be_linked(mto_api, mto):
    vehicle = VehicleFactory(unit=mto.unit, status=VehicleStatus.TERMINATED)

    response = mto_api.post(links_url(vehicle), {"person": DriverFactory(unit=mto.unit).id})

    assert response.status_code == 400
    assert response.json() == {"detail": "Only active or paused vehicles can be linked."}


def test_another_units_vehicle_is_not_found_for_linking_and_history(mto_api, mto):
    other_vehicle = VehicleFactory(unit=UnitFactory())
    driver = DriverFactory(unit=mto.unit)

    assert mto_api.post(links_url(other_vehicle), {"person": driver.id}).status_code == 404
    assert mto_api.get(links_url(other_vehicle)).status_code == 404
    assert VehicleAssignment.objects.count() == 0


# --- history ---------------------------------------------------------------------------------


def test_history_lists_current_and_ended_links_newest_first(mto_api, mto, vehicle):
    first_driver = DriverFactory(unit=mto.unit, full_name="First Driver")
    first = services.assign_person(vehicle, first_driver, mto)
    services.end_assignment(first, mto)
    officer = services.assign_person(vehicle, OfficerFactory(unit=mto.unit), mto)
    second = services.assign_person(vehicle, DriverFactory(unit=mto.unit, full_name="Second Driver"), mto)
    services.assign_person(VehicleFactory(unit=mto.unit), DriverFactory(unit=mto.unit), mto)  # another vehicle

    response = mto_api.get(links_url(vehicle))

    assert response.status_code == 200
    rows = response.json()
    assert [row["id"] for row in rows] == [second.id, officer.id, first.id]
    assert rows[0]["ended_at"] is None and rows[0]["ended_by_name"] is None
    assert rows[2]["person_name"] == "First Driver"
    assert rows[2]["ended_at"] is not None
    assert rows[2]["ended_by_name"] == mto.full_name


def test_history_uses_a_fixed_number_of_queries(mto_api, mto, vehicle, django_assert_max_num_queries):
    for _ in range(3):
        link = services.assign_person(vehicle, DriverFactory(unit=mto.unit), mto)
        services.end_assignment(link, mto)

    with django_assert_max_num_queries(6):
        response = mto_api.get(links_url(vehicle))

    assert len(response.json()) == 3
    assert all(row["ended_by_name"] == mto.full_name for row in response.json())


# --- ending a link ---------------------------------------------------------------------------


def test_ending_a_link_sets_ended_at_and_the_history_names_who_ended_it(mto_api, mto, vehicle):
    driver = DriverFactory(unit=mto.unit)
    link = services.assign_person(vehicle, driver, mto)

    response = mto_api.post(f"/api/assignments/{link.id}/end/")

    assert response.status_code == 200
    body = response.json()
    assert body["id"] == link.id
    assert body["ended_at"] is not None
    assert body["ended_by_name"] == mto.full_name
    link.refresh_from_db()
    assert link.ended_at is not None and link.ended_by == mto
    assert mto_api.get(f"/api/vehicles/{vehicle.id}/").json()["current_driver"] is None

    history = mto_api.get(links_url(vehicle)).json()
    assert [(row["id"], row["ended_by_name"]) for row in history] == [(link.id, mto.full_name)]

    # The driver is free again.
    assert mto_api.post(links_url(vehicle), {"person": driver.id}).status_code == 201


def test_ending_a_link_twice_is_refused(mto_api, mto, vehicle):
    link = services.assign_person(vehicle, DriverFactory(unit=mto.unit), mto)
    assert mto_api.post(f"/api/assignments/{link.id}/end/").status_code == 200

    again = mto_api.post(f"/api/assignments/{link.id}/end/")

    assert again.status_code == 400
    assert again.json() == {"detail": "This link has already ended."}


def test_ending_another_units_link_is_not_found(mto_api, mto):
    other = MTOFactory()
    other_vehicle = VehicleFactory(unit=other.unit)
    link = services.assign_person(other_vehicle, DriverFactory(unit=other.unit), other)

    response = mto_api.post(f"/api/assignments/{link.id}/end/")

    assert response.status_code == 404
    link.refresh_from_db()
    assert link.ended_at is None


def test_assignments_only_offer_the_end_action(mto_api, mto, vehicle):
    link = services.assign_person(vehicle, DriverFactory(unit=mto.unit), mto)

    assert mto_api.get("/api/assignments/").status_code == 404
    assert mto_api.get(f"/api/assignments/{link.id}/").status_code == 404
    assert mto_api.delete(f"/api/assignments/{link.id}/").status_code == 404
    assert mto_api.get(f"/api/assignments/{link.id}/end/").status_code == 405


@pytest.mark.parametrize("make_user", [DriverFactory, OfficerFactory, PTOFactory])
def test_only_the_mto_can_manage_links(api, mto, vehicle, make_user):
    link = services.assign_person(vehicle, DriverFactory(unit=mto.unit), mto)
    kwargs = {} if make_user is PTOFactory else {"unit": mto.unit}
    api.force_login(make_user(**kwargs))

    assert api.post(links_url(vehicle), {"person": link.person_id}).status_code == 403
    assert api.get(links_url(vehicle)).status_code == 403
    assert api.post(f"/api/assignments/{link.id}/end/").status_code == 403
    link.refresh_from_db()
    assert link.ended_at is None


def test_anonymous_is_refused(api, mto, vehicle):
    link = services.assign_person(vehicle, DriverFactory(unit=mto.unit), mto)

    assert api.get(links_url(vehicle)).status_code in (401, 403)
    assert api.post(f"/api/assignments/{link.id}/end/").status_code in (401, 403)


# --- my vehicles -----------------------------------------------------------------------------


def test_an_officer_sees_only_their_vehicles_with_the_driver_and_the_mto_contact(api, mto):
    officer = OfficerFactory(unit=mto.unit)
    mine_b = VehicleFactory(unit=mto.unit, registration_number="AP39ZZ0002")
    mine_a = VehicleFactory(unit=mto.unit, registration_number="AP39ZZ0001")
    ended = VehicleFactory(unit=mto.unit)
    not_mine = VehicleFactory(unit=mto.unit)
    driver = DriverFactory(unit=mto.unit, full_name="Driver One", mobile="9876500002")
    services.assign_person(mine_b, officer, mto)
    services.assign_person(mine_a, officer, mto)
    services.assign_person(mine_a, driver, mto)
    services.end_assignment(services.assign_person(ended, officer, mto), mto)
    services.assign_person(not_mine, OfficerFactory(unit=mto.unit), mto)
    api.force_login(officer)

    response = api.get("/api/me/vehicles/")

    assert response.status_code == 200
    body = response.json()
    assert set(body) == {"vehicles", "mto"}
    assert [v["registration_number"] for v in body["vehicles"]] == ["AP39ZZ0001", "AP39ZZ0002"]
    first = body["vehicles"][0]
    assert first["current_officer"]["id"] == officer.id
    assert first["current_driver"]["full_name"] == "Driver One"
    assert first["current_driver"]["mobile"] == "9876500002"
    assert body["vehicles"][1]["current_driver"] is None
    assert body["mto"] == {"unit_name": mto.unit.name, "full_name": "Sri K. Ramesh", "mobile": "9876500001"}


def test_a_driver_sees_their_vehicle_with_the_officer(api, mto, vehicle):
    driver = DriverFactory(unit=mto.unit)
    officer = OfficerFactory(unit=mto.unit)
    services.assign_person(vehicle, driver, mto)
    services.assign_person(vehicle, officer, mto)
    services.assign_person(VehicleFactory(unit=mto.unit), DriverFactory(unit=mto.unit), mto)
    api.force_login(driver)

    response = api.get("/api/me/vehicles/")

    assert response.status_code == 200
    body = response.json()
    assert [v["id"] for v in body["vehicles"]] == [vehicle.id]
    assert body["vehicles"][0]["current_officer"]["id"] == officer.id
    assert body["vehicles"][0]["current_driver"]["id"] == driver.id
    assert body["mto"]["unit_name"] == mto.unit.name


def test_each_vehicle_says_the_lowest_odometer_reading_it_can_take_next(api, mto):
    """The highest figure on record: the onboarding reading, the newest Sunday reading or the newest service's."""
    fresh = VehicleFactory(unit=mto.unit, odometer_at_onboarding_km=12000)
    read = VehicleFactory(unit=mto.unit, odometer_at_onboarding_km=12000)
    serviced = VehicleFactory(unit=mto.unit, odometer_at_onboarding_km=12000)
    OdometerReadingFactory(vehicle=read, reading_km=12500)
    OdometerReadingFactory(vehicle=serviced, reading_km=12500)
    ServiceRecordFactory(vehicle=serviced, odometer_km=13100)
    officer = OfficerFactory(unit=mto.unit)
    for vehicle in (fresh, read, serviced):
        services.assign_person(vehicle, officer, mto)
    api.force_login(officer)

    vehicles = api.get("/api/me/vehicles/").json()["vehicles"]

    assert {v["id"]: v["latest_odometer_km"] for v in vehicles} == {fresh.id: 12000, read.id: 12500, serviced.id: 13100}


def test_a_person_with_no_vehicle_gets_an_empty_list_and_the_contact(api, mto):
    driver = DriverFactory(unit=mto.unit)
    api.force_login(driver)

    body = api.get("/api/me/vehicles/").json()

    assert body["vehicles"] == []
    assert body["mto"]["full_name"] == "Sri K. Ramesh"


def test_a_unit_without_an_mto_gives_a_null_contact(api):
    driver = DriverFactory()  # its unit has no MTO login
    api.force_login(driver)

    body = api.get("/api/me/vehicles/").json()

    assert body["mto"] == {"unit_name": driver.unit.name, "full_name": None, "mobile": None}


def test_my_vehicles_uses_a_fixed_number_of_queries(api, mto, django_assert_max_num_queries):
    officer = OfficerFactory(unit=mto.unit)
    for _ in range(4):
        vehicle = VehicleFactory(unit=mto.unit)
        services.assign_person(vehicle, officer, mto)
        services.assign_person(vehicle, DriverFactory(unit=mto.unit), mto)
    api.force_login(officer)

    with django_assert_max_num_queries(8):
        response = api.get("/api/me/vehicles/")

    assert all(v["current_driver"] for v in response.json()["vehicles"])


@pytest.mark.parametrize("make_user", [MTOFactory, PTOFactory])
def test_other_roles_cannot_use_my_vehicles(api, make_user):
    api.force_login(make_user())

    assert api.get("/api/me/vehicles/").status_code == 403


def test_my_vehicles_needs_a_login(api):
    assert api.get("/api/me/vehicles/").status_code in (401, 403)


# --- the link history as an Excel file -------------------------------------------------------


def test_the_link_history_downloads_with_drivers_and_officers_on_their_own_sheets(mto_api, mto, vehicle):
    from datetime import datetime
    from io import BytesIO

    from openpyxl import load_workbook

    from common.months import IST

    vehicle.registration_number = "AP39PA1001"
    vehicle.save()
    first = DriverFactory(unit=mto.unit, full_name="Shaik Imran", emp_id="AP4002")
    second = DriverFactory(unit=mto.unit, full_name="Ravi Kumar", emp_id="AP4001")
    officer = OfficerFactory(unit=mto.unit, full_name="S. Venkata Rao", emp_id="AP3001")
    ended = services.assign_person(vehicle, first, mto)
    services.end_assignment(ended, mto)
    services.assign_person(vehicle, second, mto)
    services.assign_person(vehicle, officer, mto)
    VehicleAssignment.objects.filter(person=first).update(
        started_at=datetime(2026, 8, 1, 9, 0, tzinfo=IST), ended_at=datetime(2026, 9, 1, 9, 0, tzinfo=IST)
    )
    VehicleAssignment.objects.filter(person__in=[second, officer]).update(
        started_at=datetime(2026, 9, 1, 10, 0, tzinfo=IST)
    )

    response = mto_api.get(f"{links_url(vehicle)}export/")

    assert response.status_code == 200
    assert response["Content-Disposition"] == 'attachment; filename="link-history-AP39PA1001.xlsx"'
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Drivers", "Officers"]
    assert [[cell.value for cell in row] for row in book["Drivers"].iter_rows()] == [
        ["Name", "Emp ID", "From", "To", "Linked by", "Ended by"],
        ["Ravi Kumar", "AP4001", datetime(2026, 9, 1, 10, 0), "Current", mto.full_name, None],
        ["Shaik Imran", "AP4002", datetime(2026, 8, 1, 9, 0), datetime(2026, 9, 1, 9, 0), mto.full_name, mto.full_name],
    ]
    assert [[cell.value for cell in row] for row in book["Officers"].iter_rows()][1][:2] == ["S. Venkata Rao", "AP3001"]


def test_another_offices_link_history_cannot_be_downloaded(api, vehicle):
    api.force_login(MTOFactory())

    assert api.get(f"{links_url(vehicle)}export/").status_code == 404
