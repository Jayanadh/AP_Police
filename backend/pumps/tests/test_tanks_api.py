from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook

from fleet.models import FuelType
from notifications.models import Notification
from pumps import stock
from pumps.models import Pump, PumpKind, StockEntry, StockEntryKind
from testing.factories import (
    DriverFactory,
    MTOFactory,
    OfficerFactory,
    PTOFactory,
    PumpStaffFactory,
    UnitFactory,
    police_pump,
    tieup_pump,
)

pytestmark = pytest.mark.django_db


@pytest.fixture
def mto():
    return MTOFactory()


@pytest.fixture
def pump(mto):
    return police_pump(mto.unit, petrol=Decimal("500"), diesel=Decimal("300"))


@pytest.fixture
def staff(pump):
    return PumpStaffFactory(pump=pump, unit=pump.unit)


@pytest.fixture
def diesel(pump):
    return pump.tanks.get(fuel_type=FuelType.DIESEL)


@pytest.fixture
def staff_api(api, staff):
    api.force_login(staff)
    return api


@pytest.fixture
def mto_api(api, mto):
    api.force_login(mto)
    return api


def test_staff_list_the_tanks_of_their_own_pump(staff_api, pump):
    police_pump(pump.unit)  # another pump of the same office
    response = staff_api.get("/api/tanks/")

    assert response.status_code == 200
    rows = response.json()
    assert [row["fuel_type"] for row in rows] == ["PETROL", "DIESEL"]
    row = rows[1]
    assert set(row) == {
        "id", "pump", "pump_name", "fuel_type", "current_stock_litres", "low_stock_threshold_litres",
        "capacity_litres", "is_low", "opening_set",
    }
    assert row["pump"] == pump.id
    assert row["pump_name"] == pump.name
    assert row["current_stock_litres"] == "300.00"
    assert row["low_stock_threshold_litres"] == "100.00"
    assert row["capacity_litres"] is None
    assert row["is_low"] is False
    assert row["opening_set"] is False  # nothing has been recorded for the tank yet


def test_the_tanks_come_pump_by_pump_in_a_fixed_order(mto_api, mto):
    first, second = police_pump(mto.unit), police_pump(mto.unit)

    rows = mto_api.get("/api/tanks/").json()

    expected = [(pump.id, tank.id) for pump in (first, second) for tank in pump.tanks.order_by("id")]
    assert [(row["pump"], row["id"]) for row in rows] == expected


def test_the_mto_lists_the_tanks_of_all_the_pumps_of_their_unit(mto_api, mto, pump):
    second = police_pump(mto.unit)
    tieup_pump(mto.unit)
    police_pump(UnitFactory())

    ids = {row["pump"] for row in mto_api.get("/api/tanks/").json()}

    assert ids == {pump.id, second.id}


def test_staff_of_a_tie_up_bunk_see_no_tanks(api, mto):
    bunk = tieup_pump(mto.unit)
    api.force_login(PumpStaffFactory(pump=bunk, unit=mto.unit))
    response = api.get("/api/tanks/")
    assert response.status_code == 200
    assert response.json() == []


def test_a_pump_switched_to_a_bunk_no_longer_shows_its_old_tanks(api, mto, staff, pump):
    Pump.objects.filter(pk=pump.pk).update(kind=PumpKind.TIE_UP)  # the tanks stay in the database, as history
    assert pump.tanks.count() == 2
    for person in (staff, mto):
        api.force_login(person)
        assert api.get("/api/tanks/").json() == []


def test_staff_cannot_open_a_tank_of_another_pump(staff_api, mto):
    other = police_pump(mto.unit).tanks.first()
    assert staff_api.get(f"/api/tanks/{other.id}/").status_code == 404
    assert staff_api.post(f"/api/tanks/{other.id}/receive/", {"litres": "10"}).status_code == 404


def test_another_units_tank_is_a_404_for_the_mto(mto_api):
    foreign = police_pump(UnitFactory()).tanks.first()
    assert mto_api.get(f"/api/tanks/{foreign.id}/").status_code == 404
    assert mto_api.patch(f"/api/tanks/{foreign.id}/", {"low_stock_threshold_litres": "50"}).status_code == 404
    assert mto_api.get(f"/api/tanks/{foreign.id}/entries/").status_code == 404


@pytest.mark.parametrize("factory", [PTOFactory, OfficerFactory, DriverFactory])
def test_other_roles_are_forbidden(api, diesel, factory):
    api.force_login(factory())
    assert api.get("/api/tanks/").status_code == 403
    assert api.get(f"/api/tanks/{diesel.id}/").status_code == 403
    assert api.get(f"/api/tanks/{diesel.id}/entries/").status_code == 403


def test_login_is_required(api, diesel):
    assert api.get("/api/tanks/").status_code in (401, 403)


def test_the_mto_sets_a_tanks_opening_stock(mto_api, mto, diesel):
    response = mto_api.post(f"/api/tanks/{diesel.id}/opening/", {"litres": "280.50", "note": "Dip stick"})

    assert response.status_code == 200
    body = response.json()
    assert body["id"] == diesel.id
    assert body["current_stock_litres"] == "280.50"
    assert body["opening_set"] is True
    entry = StockEntry.objects.get()
    assert (entry.kind, entry.litres, entry.stock_before, entry.stock_after) == (
        StockEntryKind.OPENING, Decimal("280.50"), Decimal("300"), Decimal("280.50"),
    )
    assert entry.note == "Dip stick"
    assert entry.recorded_by == mto


def test_the_opening_stock_cannot_be_set_after_a_receipt(api, mto, staff, diesel):
    api.force_login(staff)
    api.post(f"/api/tanks/{diesel.id}/receive/", {"litres": "100"})
    api.force_login(mto)

    response = api.post(f"/api/tanks/{diesel.id}/opening/", {"litres": "50"})

    assert response.status_code == 400
    assert response.json() == {
        "detail": "This tank's opening stock is already set. Its stock now changes only with tanker receipts and fills."
    }
    diesel.refresh_from_db()
    assert diesel.current_stock_litres == Decimal("400")


def test_staff_record_a_tanker_receipt(staff_api, diesel):
    response = staff_api.post(f"/api/tanks/{diesel.id}/receive/", {"litres": "1000"})

    assert response.status_code == 200
    body = response.json()
    assert body["current_stock_litres"] == "1300.00"
    assert body["opening_set"] is True
    entry = StockEntry.objects.get()
    assert entry.kind == StockEntryKind.TANKER_RECEIPT
    assert entry.note == ""


def test_only_the_mto_sets_the_opening_stock_and_only_staff_record_receipts(api, mto, staff, diesel):
    api.force_login(staff)
    assert api.post(f"/api/tanks/{diesel.id}/opening/", {"litres": "10"}).status_code == 403
    api.force_login(mto)
    assert api.post(f"/api/tanks/{diesel.id}/receive/", {"litres": "10"}).status_code == 403
    diesel.refresh_from_db()
    assert diesel.current_stock_litres == Decimal("300")
    assert StockEntry.objects.count() == 0


def test_there_is_no_daily_measurement_any_more(api, mto, staff, diesel):
    for person in (staff, mto):
        api.force_login(person)
        assert api.post(f"/api/tanks/{diesel.id}/measure/", {"litres": "10"}).status_code == 404


def test_another_units_mto_cannot_set_the_opening_stock(api, diesel):
    api.force_login(MTOFactory())
    assert api.post(f"/api/tanks/{diesel.id}/opening/", {"litres": "10"}).status_code == 404


def test_a_negative_opening_stock_is_a_400(mto_api, diesel):
    response = mto_api.post(f"/api/tanks/{diesel.id}/opening/", {"litres": "-1"})
    assert response.status_code == 400
    assert response.json() == {"detail": "Stock can't be negative."}


def test_an_empty_receipt_is_a_400(staff_api, diesel):
    response = staff_api.post(f"/api/tanks/{diesel.id}/receive/", {"litres": "0"})
    assert response.status_code == 400
    assert response.json() == {"detail": "Enter the litres received."}


@pytest.mark.parametrize("who, action", [("mto", "opening"), ("staff", "receive")])
@pytest.mark.parametrize("payload", [{}, {"litres": ""}, {"litres": "abc"}, {"litres": "1.234"}])
def test_litres_must_be_a_number_with_two_decimals(api, mto, staff, diesel, who, action, payload):
    api.force_login(mto if who == "mto" else staff)
    response = api.post(f"/api/tanks/{diesel.id}/{action}/", payload)
    assert response.status_code == 400
    assert "litres" in response.json()
    assert StockEntry.objects.count() == 0


def test_an_opening_stock_below_the_alert_level_alerts_the_mto(mto_api, mto, diesel):
    mto_api.post(f"/api/tanks/{diesel.id}/opening/", {"litres": "80"})

    body = mto_api.get(f"/api/tanks/{diesel.id}/").json()
    assert body["is_low"] is True
    assert Notification.objects.filter(recipient=mto, link=f"/mto/pumps/{diesel.pump_id}").count() == 1


def test_the_entries_of_a_month_are_listed_newest_first(staff_api, staff, diesel):
    january = timezone.now().replace(year=2026, month=1, day=15, hour=10)
    february_early = timezone.now().replace(year=2026, month=2, day=3, hour=9)
    february_late = timezone.now().replace(year=2026, month=2, day=20, hour=9)
    first = stock.record_opening(diesel, Decimal("280"), staff)
    second = stock.record_receipt(diesel, Decimal("20"), staff, note="Tanker")
    third = stock.dispense(diesel, Decimal("30"), staff)
    StockEntry.objects.filter(pk=first.pk).update(recorded_at=january)
    StockEntry.objects.filter(pk=second.pk).update(recorded_at=february_early)
    StockEntry.objects.filter(pk=third.pk).update(recorded_at=february_late)

    response = staff_api.get(f"/api/tanks/{diesel.id}/entries/?month=2026-02")

    assert response.status_code == 200
    rows = response.json()
    assert [row["id"] for row in rows] == [third.id, second.id]
    assert set(rows[0]) == {
        "id", "kind", "kind_label", "litres", "stock_before", "stock_after", "note", "recorded_by_name",
        "recorded_at",
    }
    assert rows[0]["kind"] == "DISPENSE"
    assert rows[0]["kind_label"] == "Fill"
    assert rows[0]["litres"] == "30.00"
    assert (rows[0]["stock_before"], rows[0]["stock_after"]) == ("300.00", "270.00")
    assert rows[1]["kind_label"] == "Tanker receipt"
    assert rows[1]["note"] == "Tanker"
    assert rows[0]["recorded_by_name"] == staff.full_name


def test_entries_use_asia_kolkata_month_boundaries(staff_api, staff, diesel):
    entry = stock.record_opening(diesel, Decimal("280"), staff)
    # 20:00 UTC on 30 September is 01:30 on 1 October in Kolkata.
    StockEntry.objects.filter(pk=entry.pk).update(recorded_at="2026-09-30T20:00:00+00:00")

    october = staff_api.get(f"/api/tanks/{diesel.id}/entries/?month=2026-10").json()
    september = staff_api.get(f"/api/tanks/{diesel.id}/entries/?month=2026-09").json()

    assert [row["id"] for row in october] == [entry.id]
    assert september == []


def test_entries_default_to_the_current_month(staff_api, staff, diesel):
    entry = stock.record_opening(diesel, Decimal("280"), staff)
    old = stock.record_receipt(diesel, Decimal("5"), staff)
    StockEntry.objects.filter(pk=old.pk).update(recorded_at="2020-01-10T10:00:00+00:00")

    rows = staff_api.get(f"/api/tanks/{diesel.id}/entries/").json()

    assert [row["id"] for row in rows] == [entry.id]


@pytest.mark.parametrize("month", ["2026-13", "9999-12", "1999-12"])
def test_entries_refuse_a_bad_month(staff_api, diesel, month):
    response = staff_api.get(f"/api/tanks/{diesel.id}/entries/?month={month}")
    assert response.status_code == 400
    assert response.json() == {"detail": "Use the month format YYYY-MM."}


def test_the_entries_of_any_period_are_listed(staff_api, staff, diesel):
    receipts = [stock.record_receipt(diesel, Decimal(n), staff) for n in ("5", "6", "7")]
    for entry, day in zip(receipts, ("2026-04-01", "2026-09-30", "2027-03-31")):
        StockEntry.objects.filter(pk=entry.pk).update(recorded_at=f"{day}T10:00:00+05:30")

    year = staff_api.get(f"/api/tanks/{diesel.id}/entries/", {"from": "2026-04-01", "to": "2026-09-30"}).json()

    assert [row["id"] for row in year] == [receipts[1].id, receipts[0].id]


def test_a_bad_period_is_refused(staff_api, diesel):
    response = staff_api.get(f"/api/tanks/{diesel.id}/entries/", {"from": "2026-04-01", "to": "2027-05-01"})
    assert response.status_code == 400
    assert response.json() == {"detail": "Pick a period of at most a year."}


def test_pump_staff_download_the_stock_entries_of_a_period_as_excel(staff_api, staff, pump, diesel):
    petrol = pump.tanks.get(fuel_type=FuelType.PETROL)
    opening = stock.record_opening(diesel, Decimal("280"), staff)
    receipt = stock.record_receipt(diesel, Decimal("1500.5"), staff, note="Tanker TN-1")
    StockEntry.objects.filter(pk=opening.pk).update(recorded_at="2026-09-01T09:00:00+05:30")
    StockEntry.objects.filter(pk=receipt.pk).update(recorded_at="2026-09-02T10:30:00+05:30")
    stock.record_opening(petrol, Decimal("100"), staff)  # in October: outside the period

    response = staff_api.get("/api/tanks/entries/export/", {"from": "2026-09-01", "to": "2026-09-30"})

    assert response.status_code == 200
    assert response["Content-Disposition"] == (
        f'attachment; filename="stock-{pump.name.lower().replace(" ", "-")}-2026-09-01-to-2026-09-30.xlsx"'
    )
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Petrol", "Diesel"]
    headers = ["When", "Entry", "Litres", "Stock before (L)", "Stock after (L)", "Note", "By"]
    assert [[cell.value for cell in row] for row in book["Petrol"].iter_rows()] == [headers]
    assert [[cell.value for cell in row] for row in book["Diesel"].iter_rows()] == [
        headers,
        [datetime(2026, 9, 1, 9, 0), "Opening stock", 280, 300, 280, None, staff.full_name],
        [datetime(2026, 9, 2, 10, 30), "Tanker receipt", 1500.5, 280, 1780.5, "Tanker TN-1", staff.full_name],
    ]


@pytest.mark.parametrize("maker", [MTOFactory, DriverFactory, OfficerFactory, PTOFactory])
def test_only_pump_staff_download_their_stock_entries(api, maker):
    api.force_login(maker())
    assert api.get("/api/tanks/entries/export/").status_code == 403


def test_the_mto_can_read_the_entries_of_their_pump(mto_api, staff, diesel):
    stock.record_opening(diesel, Decimal("280"), staff)
    assert len(mto_api.get(f"/api/tanks/{diesel.id}/entries/").json()) == 1


def test_the_mto_updates_the_threshold_and_capacity(mto_api, diesel):
    response = mto_api.patch(
        f"/api/tanks/{diesel.id}/", {"low_stock_threshold_litres": "150", "capacity_litres": "5000"}
    )

    assert response.status_code == 200
    body = response.json()
    assert body["low_stock_threshold_litres"] == "150.00"
    assert body["capacity_litres"] == "5000.00"
    diesel.refresh_from_db()
    assert diesel.low_stock_threshold_litres == Decimal("150")
    assert diesel.capacity_litres == Decimal("5000")
    assert body["is_low"] is False  # 300 L against a 150 L alert level


def test_lowering_the_threshold_re_arms_the_alert(mto_api, mto, diesel, staff):
    stock.dispense(diesel, Decimal("250"), staff)  # 50 L against a 100 L alert level: alerted
    assert Notification.objects.filter(recipient=mto).count() == 1

    mto_api.patch(f"/api/tanks/{diesel.id}/", {"low_stock_threshold_litres": "20"})  # 50 L is now fine

    diesel.refresh_from_db()
    assert diesel.low_stock_alerted is False
    assert Notification.objects.filter(recipient=mto).count() == 1  # lowering the level alerts nobody

    stock.dispense(diesel, Decimal("40"), staff)  # 10 L, below the new 20 L level

    assert Notification.objects.filter(recipient=mto).count() == 2
    assert Notification.objects.filter(recipient=staff).count() == 2
    newest = Notification.objects.filter(recipient=mto).first()
    assert newest.body == "10.00 L left (alert level 20.00 L)."


def test_raising_the_threshold_above_the_stock_alerts_once(mto_api, mto, staff, pump, diesel):
    response = mto_api.patch(f"/api/tanks/{diesel.id}/", {"low_stock_threshold_litres": "400"})  # 300 L in stock

    assert response.status_code == 200
    assert response.json()["is_low"] is True
    diesel.refresh_from_db()
    assert diesel.low_stock_alerted is True
    mto_alert = Notification.objects.get(recipient=mto)
    assert mto_alert.title == f"Low diesel stock at {pump.name}"
    assert mto_alert.body == "300.00 L left (alert level 400.00 L)."
    assert mto_alert.link == f"/mto/pumps/{pump.id}"
    staff_alert = Notification.objects.get(recipient=staff)
    assert staff_alert.link == "/pump/stock"

    mto_api.patch(f"/api/tanks/{diesel.id}/", {"low_stock_threshold_litres": "500"})  # still low, already alerted

    assert Notification.objects.count() == 2


def test_changing_only_the_capacity_leaves_the_alert_alone(mto_api, mto, diesel):
    mto_api.patch(f"/api/tanks/{diesel.id}/", {"capacity_litres": "5000"})
    diesel.refresh_from_db()
    assert diesel.low_stock_alerted is False
    assert Notification.objects.count() == 0


def test_patch_cannot_change_the_stock(mto_api, diesel):
    response = mto_api.patch(
        f"/api/tanks/{diesel.id}/", {"current_stock_litres": "9999", "fuel_type": "PETROL", "pump": 1}
    )
    assert response.status_code == 200
    diesel.refresh_from_db()
    assert diesel.current_stock_litres == Decimal("300")
    assert diesel.fuel_type == FuelType.DIESEL


def test_a_negative_threshold_is_refused(mto_api, diesel):
    response = mto_api.patch(f"/api/tanks/{diesel.id}/", {"low_stock_threshold_litres": "-1"})
    assert response.status_code == 400
    assert "low_stock_threshold_litres" in response.json()


def test_staff_cannot_change_the_threshold(staff_api, diesel):
    response = staff_api.patch(f"/api/tanks/{diesel.id}/", {"low_stock_threshold_litres": "1"})
    assert response.status_code == 403
    diesel.refresh_from_db()
    assert diesel.low_stock_threshold_litres == Decimal("100")


def test_tanks_cannot_be_created_or_deleted_through_the_api(mto_api, diesel):
    assert mto_api.post("/api/tanks/", {"pump": diesel.pump_id, "fuel_type": "PETROL"}).status_code == 405
    assert mto_api.delete(f"/api/tanks/{diesel.id}/").status_code == 405
    assert mto_api.put(f"/api/tanks/{diesel.id}/", {}).status_code == 405


def test_the_tank_list_uses_a_fixed_number_of_queries(mto_api, mto):
    for _ in range(4):
        police_pump(mto.unit)
    with CaptureQueriesContext(connection) as queries:
        response = mto_api.get("/api/tanks/")
    assert len(response.json()) == 8
    assert len(queries) <= 4  # session, user, tanks (with pump name and whether the opening stock is set)
